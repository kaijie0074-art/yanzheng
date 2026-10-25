"""不调用外部模型：以真实PDF导入、受控模型回复检查来源边界与版本行为。"""

import hashlib
import json
from io import BytesIO

import pytest
from docx import Document
from reportlab.pdfgen.canvas import Canvas

from research_app import documents, paper_work
from research_app.backends import BackendError
from research_app.models import BriefContent, ModelConfig, ModelReply, now
from research_app.storage import Store, encoded
from research_app.workflow import confirm_brief


def confirm(store, question="数字学习环境如何影响学习效果？"):
    return confirm_brief(
        store,
        BriefContent(
            question=question,
            population_or_object="大学生",
            constraints="只用合法取得的公开资料",
            available_resources="本地资料与模型",
            validation_goal="形成待核验的研究综述",
        ),
        ["digital learning outcomes"],
    )


@pytest.fixture
def store(tmp_path):
    value = Store(tmp_path)
    confirm(value)
    return value


def abstract_source(
    store,
    text="Digital learning supports access. Context limits transfer.",
    title="Learning evidence",
):
    sid = store.add_source(
        title=title,
        authors=["Known Author"],
        year=2024,
        official_url="https://example.org/research",
        doi="10.1234/known",
        access_level="abstract",
    )
    with store.connection() as db:
        db.execute(
            "CREATE TABLE IF NOT EXISTS paper_details(source_id TEXT PRIMARY KEY REFERENCES sources(id), metadata_json TEXT NOT NULL)"
        )
        db.execute(
            "INSERT INTO paper_details VALUES(?,?)",
            (
                sid,
                encoded(
                    {
                        "abstract": text,
                        "provider": "crossref",
                        "remote_id": "known",
                        "pdf_url": None,
                        "license": None,
                    }
                ),
            ),
        )
    return sid


def pdf_source(store, pages=None):
    output = BytesIO()
    pdf = Canvas(output, invariant=1)
    for text in pages or [
        "Digital learning supports access.",
        "Context limits transfer.",
    ]:
        pdf.drawString(40, 700, text)
        pdf.showPage()
    pdf.save()
    source = documents.import_pdf(
        store,
        output.getvalue(),
        title="Full text evidence",
        authors=["Real Author"],
        year=2023,
        official_url="https://example.org/fulltext",
    )
    return source["id"]


def citation(payload):
    excerpt = payload["excerpts"][0]
    return {
        "source_id": excerpt["source_id"],
        "page_number": excerpt["page_number"],
        "quote": excerpt["text"].strip(),
    }


def reading_result(payload):
    return {
        "findings": [{**citation(payload), "text": "文献讨论了学习环境及其适用边界。"}],
        "limitations": ["原文命中不等于研究结论已被验证。"],
    }


def writing_result(payload):
    citations = [citation(payload)] if payload["excerpts"] else []
    return {
        "title": "学习环境研究草稿",
        "sections": [
            {
                "heading": "研究背景",
                "paragraphs": [
                    {"text": "文献讨论了学习环境及其适用边界。", "citations": citations}
                ],
            }
        ],
    }


class Backend:
    config = ModelConfig()

    def __init__(self, result=reading_result, mutate=None):
        self.result = result
        self.mutate = mutate
        self.calls = []

    def complete_json(self, *, system, payload, schema):
        self.calls.append({"system": system, "payload": payload})
        result = self.result(payload) if callable(self.result) else self.result
        value = schema.model_validate(result)
        if self.mutate:
            self.mutate()
        return ModelReply(
            value=value,
            requested_model="test/model",
            actual_model="test/model",
            backend="local_account",
            usage=None,
            called_at=now(),
            text=encoded(result),
        )


def read(store, sid, backend=None):
    return paper_work.read_sources(
        store,
        backend or Backend(),
        source_ids=[sid],
        question="学习环境有何适用边界？",
        rights_confirmed=True,
    )


def write_draft(store, sid, backend=None, **kwargs):
    return paper_work.generate_writing(
        store,
        backend or Backend(writing_result),
        kind=kwargs.pop("kind", "draft"),
        instructions="形成谨慎的研究背景草稿",
        source_ids=[sid],
        rights_confirmed=True,
        **kwargs,
    )


def assert_failed(store, code="CONTENT_VALIDATION"):
    runs = store.rows("SELECT * FROM model_runs ORDER BY called_at,id")
    assert runs[-1]["status"] == "failed"
    assert runs[-1]["error_type"] == code


def test_metadata_only_does_not_call_model_or_generate_discoveries(store):
    sid = store.add_source(
        title="Only metadata",
        authors=[],
        year=None,
        official_url="https://example.org/meta",
    )
    backend = Backend()
    with pytest.raises(ValueError, match="没有可读正文或摘要"):
        read(store, sid, backend)
    with pytest.raises(ValueError, match="至少需要"):
        write_draft(store, sid, backend)
    assert backend.calls == []
    assert store.rows("SELECT * FROM model_runs") == []
    assert paper_work.get_state(store) == {"reading_reports": [], "writing_drafts": []}


def test_abstract_does_not_need_bibliography_or_evidence_approval(store):
    sid = abstract_source(store)
    result = read(store, sid)
    assert result["findings"][0]["page_number"] is None
    assert result["coverage"] == [
        {
            "source_id": sid,
            "title": "Learning evidence",
            "level": "abstract",
            "pages": [],
            "total_pages": 0,
            "characters": len(
                "Digital learning supports access. Context limits transfer."
            ),
        }
    ]
    assert any("仅有摘要" in item for item in result["limitations"])
    assert store.get_source(sid)["bibliography_status"] == "unverified"
    assert store.rows("SELECT * FROM evidence_cards") == []
    assert paper_work.get_state(Store(store.data_dir))["reading_reports"] == [result]


def test_local_pdf_without_discovery_tables_has_physical_page_coverage(store):
    sid = pdf_source(store)
    backend = Backend()
    result = read(store, sid, backend)
    assert result["coverage"][0]["level"] == "fulltext"
    assert result["coverage"][0]["pages"] == [1, 2]
    assert result["coverage"][0]["total_pages"] == 2
    excerpt = backend.calls[0]["payload"]["excerpts"][0]
    assert result["findings"][0]["quote"] in excerpt["text"]
    assert result["findings"][0]["page_number"] == excerpt["page_number"]
    run = store.rows("SELECT * FROM model_runs")[0]
    refs = json.loads(run["input_refs"])
    assert refs["source_fingerprints"][sid]
    assert refs["excerpts"][0]["text_sha256"]
    assert "text" not in refs["excerpts"][0]


@pytest.mark.parametrize(
    "alter",
    [
        lambda c: c.update(source_id="invented"),
        lambda c: c.update(page_number=777),
        lambda c: c.update(quote="A sentence that was never supplied."),
        lambda c: c.update(quote=" "),
    ],
)
def test_unknown_sources_pages_and_false_quotes_are_failed_audits(store, alter):
    sid = pdf_source(store)

    def result(payload):
        data = reading_result(payload)
        alter(data["findings"][0])
        return data

    with pytest.raises(ValueError, match="本次实际发送"):
        read(store, sid, Backend(result))
    assert_failed(store)
    assert paper_work.get_state(store)["reading_reports"] == []
    assert store.get_source(sid)["access_level"] == "fulltext"


@pytest.mark.parametrize("operation", ["read", "write"])
@pytest.mark.parametrize("level", ["abstract", "fulltext"])
def test_failed_citations_retain_only_actual_input_locations_and_hashes(
    store, operation, level
):
    sid = (
        abstract_source(store, "Digital learning supports access. " * 100)
        if level == "abstract"
        else pdf_source(store)
    )
    skipped = store.add_source(
        title="Unread metadata",
        authors=[],
        year=None,
        official_url="https://example.org/unread",
    )

    def result(payload):
        data = (reading_result if operation == "read" else writing_result)(payload)
        target = (
            data["findings"][0]
            if operation == "read"
            else data["sections"][0]["paragraphs"][0]["citations"][0]
        )
        target["quote"] = "This sentence was not supplied."
        return data

    backend = Backend(result)
    with pytest.raises(ValueError):
        if operation == "read":
            paper_work.read_sources(
                store,
                backend,
                source_ids=[skipped, sid],
                question="learning",
                rights_confirmed=True,
            )
        else:
            paper_work.generate_writing(
                store,
                backend,
                kind="draft",
                instructions="讨论学习环境",
                source_ids=[skipped, sid],
                rights_confirmed=True,
            )
    assert_failed(store)
    assert paper_work.get_state(store) == {"reading_reports": [], "writing_drafts": []}
    excerpts = backend.calls[0]["payload"]["excerpts"]
    run = store.rows("SELECT * FROM model_runs")[0]
    refs = json.loads(run["input_refs"])
    assert refs["source_ids"] == [sid]
    expected_pages = [None] if level == "abstract" else [1, 2]
    assert refs["pages"] == [
        {"source_id": sid, "page_number": page} for page in expected_pages
    ]
    assert set(refs["source_fingerprints"]) == {sid}
    assert len(refs["source_fingerprints"][sid]) == 64
    assert {item["source_id"] for item in refs["coverage"]} == {sid}
    assert refs["coverage"][0]["characters"] == sum(
        len(item["text"]) for item in excerpts
    )
    assert refs["coverage"][0]["pages"] == ([] if level == "abstract" else [1, 2])
    assert refs["excerpts"] == [
        {key: value for key, value in item.items() if key != "text"}
        | {"text_sha256": hashlib.sha256(item["text"].encode("utf-8")).hexdigest()}
        for item in excerpts
    ]
    assert skipped not in run["input_refs"]
    assert all(item["text"] not in run["input_refs"] for item in excerpts)


def test_abstract_cannot_be_assigned_a_fake_page(store):
    sid = abstract_source(store)

    def result(payload):
        data = reading_result(payload)
        data["findings"][0]["page_number"] = 1
        return data

    with pytest.raises(ValueError):
        read(store, sid, Backend(result))
    assert_failed(store)


def test_quote_normalization_is_limited_to_nfkc_and_whitespace(store):
    sid = abstract_source(store, "Ａ  learning\n evidence")

    def result(payload):
        data = reading_result(payload)
        data["findings"][0]["quote"] = "A learning evidence"
        return data

    assert read(store, sid, Backend(result))["findings"]

    def lower_case(payload):
        data = result(payload)
        data["findings"][0]["quote"] = "a learning evidence"
        return data

    with pytest.raises(ValueError):
        read(store, sid, Backend(lower_case))
    assert_failed(store)


@pytest.mark.parametrize("level", ["abstract", "fulltext"])
def test_continuous_quote_across_sent_chunk_boundary_is_saved(store, level):
    quote = "A continuous original sentence crosses the boundary."
    text = "x" * 1590 + quote + "y" * 1700
    sid = (
        abstract_source(store, text)
        if level == "abstract"
        else pdf_source(store, [text])
    )
    backend = Backend(
        {
            "findings": [
                {
                    "source_id": sid,
                    "page_number": None if level == "abstract" else 1,
                    "quote": quote,
                    "text": "这条摘录来自实际发送的连续原文，仍待人工核验。",
                }
            ],
            "limitations": [],
        }
    )
    report = read(store, sid, backend)
    persisted = paper_work.get_state(Store(store.data_dir))["reading_reports"]
    assert persisted[0]["findings"][0]["quote"] == quote
    assert report["coverage"][0]["pages"] == ([] if level == "abstract" else [1])
    assert store.get_source(sid)["bibliography_status"] == "unverified"


def test_quote_cannot_join_sent_chunks_across_an_unsent_gap(store):
    left, right = "Left-boundary.", "Right-boundary."
    blocks = ["x" * 1600 for _ in range(17)]
    blocks[0] = "learning " + "x" * (1600 - len("learning "))
    blocks[13] = "x" * (1600 - len(left)) + left
    blocks[16] = right + " learning " + "x" * (1600 - len(right + " learning "))
    sid = abstract_source(store, "".join(blocks))
    backend = Backend(
        {
            "findings": [
                {
                    "source_id": sid,
                    "page_number": None,
                    "quote": left + right,
                    "text": "这是一条跨越未发送区间的错误拼接。",
                }
            ],
            "limitations": [],
        }
    )
    with pytest.raises(ValueError):
        paper_work.read_sources(
            store,
            backend,
            source_ids=[sid],
            question="learning",
            rights_confirmed=True,
        )
    sent = "\n".join(item["text"] for item in backend.calls[0]["payload"]["excerpts"])
    assert left in sent and right in sent
    assert paper_work.get_state(store)["reading_reports"] == []
    assert_failed(store)


def test_automatic_selection_is_bounded_and_cannot_quote_unsent_text(store):
    # 主题相关片段在末尾；大量无关正文不能挤掉它，覆盖必须实报截取量。
    unsent = "A unique observation in the unsent middle."
    text = "unrelated " * 4000 + unsent + "unrelated " * 4000
    text += " digital learning unique finding "
    sid = abstract_source(store, text)
    backend = Backend()
    result = paper_work.read_sources(
        store,
        backend,
        source_ids=[sid],
        question="digital learning",
        rights_confirmed=True,
    )
    excerpts = backend.calls[0]["payload"]["excerpts"]
    assert sum(len(item["text"]) for item in excerpts) == 24000
    assert result["coverage"][0]["characters"] == 24000
    assert any("unique finding" in item["text"] for item in excerpts)
    assert result["coverage"][0]["characters"] < len(text)

    def fabricated(payload):
        data = reading_result(payload)
        data["findings"][0]["quote"] = unsent
        return data

    # 原文确实包含此句，但它位于未发送部分，不能被允许引用。
    with pytest.raises(ValueError):
        read(store, sid, Backend(fabricated))
    assert_failed(store)


def test_each_selected_readable_source_receives_disclosed_budget(store):
    ids = [
        abstract_source(store, "learning " * 6000, title=f"Source {index}")
        for index in range(3)
    ]
    backend = Backend()
    report = paper_work.read_sources(
        store, backend, source_ids=ids, question="learning", rights_confirmed=True
    )
    assert {item["source_id"] for item in report["coverage"]} == set(ids)
    assert [item["characters"] for item in report["coverage"]] == [8000, 8000, 8000]
    assert (
        sum(len(item["text"]) for item in backend.calls[0]["payload"]["excerpts"])
        == 24000
    )


@pytest.mark.parametrize("operation", ["read", "write"])
def test_research_revision_changed_during_call_is_not_saved(store, operation):
    sid = abstract_source(store)
    backend = Backend(
        reading_result if operation == "read" else writing_result,
        mutate=lambda: confirm(store, "更换研究问题"),
    )
    with pytest.raises(paper_work.PaperWorkConflict, match="研究问题已修订"):
        (read if operation == "read" else write_draft)(store, sid, backend)
    assert_failed(store, "INPUT_CHANGED")
    assert paper_work.get_state(store) == {"reading_reports": [], "writing_drafts": []}


@pytest.mark.parametrize("change", ["pdf", "page", "abstract", "metadata"])
@pytest.mark.parametrize("operation", ["read", "write"])
def test_input_integrity_changes_during_model_call_fail_closed(
    store, change, operation
):
    sid = abstract_source(store) if change == "abstract" else pdf_source(store)

    def mutate():
        if change == "pdf":
            digest = store.get_source(sid)["sha256"]
            (store.papers_dir / f"{digest}.pdf").write_bytes(b"tampered")
        else:
            with store.connection() as db:
                if change == "page":
                    db.execute(
                        "UPDATE pages SET text='changed' WHERE source_id=?", (sid,)
                    )
                elif change == "abstract":
                    db.execute(
                        "UPDATE paper_details SET metadata_json=? WHERE source_id=?",
                        (encoded({"abstract": "changed"}), sid),
                    )
                else:
                    db.execute("UPDATE sources SET title='changed' WHERE id=?", (sid,))

    backend = Backend(reading_result if operation == "read" else writing_result, mutate)
    with pytest.raises(paper_work.PaperWorkConflict):
        (read if operation == "read" else write_draft)(store, sid, backend)
    assert_failed(store, "INPUT_CHANGED")
    assert paper_work.get_state(store) == {"reading_reports": [], "writing_drafts": []}


def test_bad_pdf_before_reading_never_calls_model(store):
    sid = pdf_source(store)
    source = store.get_source(sid)
    (store.papers_dir / f"{source['sha256']}.pdf").write_bytes(b"modified")
    backend = Backend()
    with pytest.raises(ValueError, match="哈希不一致"):
        read(store, sid, backend)
    assert backend.calls == []


def test_backend_failure_is_audited_without_losing_sources(store):
    sid = abstract_source(store)

    class Failing:
        config = ModelConfig()

        def complete_json(self, **kwargs):
            raise BackendError("RATE_LIMIT", "暂不可用")

    with pytest.raises(BackendError):
        read(store, sid, Failing())
    assert_failed(store, "RATE_LIMIT")
    assert store.get_source(sid)["title"] == "Learning evidence"
    assert paper_work.get_state(store)["reading_reports"] == []


@pytest.mark.parametrize("source_ids,rights", [([], True), (["missing"], True)])
def test_no_or_unknown_reading_sources_do_not_call_model(store, source_ids, rights):
    backend = Backend()
    with pytest.raises(ValueError):
        paper_work.read_sources(
            store,
            backend,
            source_ids=source_ids,
            question="问题",
            rights_confirmed=rights,
        )
    assert backend.calls == []


def test_explicit_rights_and_source_limit_before_paid_call(store):
    sid = abstract_source(store)
    backend = Backend()
    for ids, rights in [([sid], False), ([sid] * 9, True)]:
        with pytest.raises(ValueError):
            paper_work.read_sources(
                store, backend, source_ids=ids, question="问题", rights_confirmed=rights
            )
    assert backend.calls == []


def test_outline_without_sources_is_persisted_but_has_no_invented_references(store):
    result = paper_work.generate_writing(
        store,
        Backend(writing_result),
        kind="outline",
        instructions="给出综述结构",
        source_ids=[],
        rights_confirmed=False,
    )
    assert result["references"] == [] and result["coverage"] == []
    assert "结构建议，无文献支持" in result["markdown"]
    assert paper_work.get_state(Store(store.data_dir))["writing_drafts"] == [result]


def test_draft_citations_and_bibliography_come_from_real_unverified_source(store):
    sid = abstract_source(store)
    draft = write_draft(store, sid)
    assert draft["references"] == [
        {
            "number": 1,
            "source_id": sid,
            "title": "Learning evidence",
            "authors": ["Known Author"],
            "year": 2024,
            "doi": "10.1234/known",
            "official_url": "https://example.org/research",
            "access_level": "abstract",
            "pages": [],
            "bibliography_status": "unverified",
        }
    ]
    assert "文献讨论了学习环境及其适用边界。[1]" in draft["markdown"]
    assert "Known Author" in draft["markdown"] and "仅摘要" in draft["markdown"]
    assert "AI辅助草稿" in draft["markdown"] and "书目待人工核验" in draft["markdown"]
    assert store.rows("SELECT * FROM decisions") == []


@pytest.mark.parametrize(
    "change",
    [
        "no_citation",
        "unknown",
        "fake_quote",
        "fake_page",
        "number",
        "doi",
        "bibliography",
    ],
)
def test_writing_refuses_ungrounded_model_content(store, change):
    sid = abstract_source(store)

    def result(payload):
        data = writing_result(payload)
        section = data["sections"][0]
        paragraph = section["paragraphs"][0]
        if change == "no_citation":
            paragraph["citations"] = []
        elif change == "unknown":
            paragraph["citations"][0]["source_id"] = "unknown"
        elif change == "fake_quote":
            paragraph["citations"][0]["quote"] = "invented"
        elif change == "fake_page":
            paragraph["citations"][0]["page_number"] = 1
        elif change == "number":
            paragraph["text"] = "本研究准确率达到99.9%。"
        elif change == "doi":
            paragraph["text"] = "依据10.9999/fake研究。"
        else:
            section["heading"] = "参考文献"
        return data

    with pytest.raises(ValueError):
        write_draft(store, sid, Backend(result))
    assert_failed(store)
    assert paper_work.get_state(store)["writing_drafts"] == []


@pytest.mark.parametrize("kind", ["draft", "revision"])
@pytest.mark.parametrize("target", ["title", "heading"])
@pytest.mark.parametrize(
    "claim",
    [
        "准确率达到99.9%",
        "1.2 准确率达到99.9%",
        "第2章 准确率达到99.9%",
        "2026 年学习效果分析",
        "100 例样本的研究结果",
        "1.2 倍性能提升",
    ],
)
def test_numeric_titles_and_headings_require_actual_cited_numbers(
    store, kind, target, claim
):
    sid = abstract_source(store)
    parent = write_draft(store, sid) if kind == "revision" else None
    existing = paper_work.get_state(store)["writing_drafts"]

    def result(payload):
        data = writing_result(payload)
        (data if target == "title" else data["sections"][0])[target] = claim
        return data

    kwargs = {"parent_id": parent["id"]} if parent else {}
    with pytest.raises(ValueError):
        write_draft(store, sid, Backend(result), kind=kind, **kwargs)
    assert_failed(store)
    assert paper_work.get_state(store)["writing_drafts"] == existing


@pytest.mark.parametrize("kind", ["draft", "revision"])
@pytest.mark.parametrize(
    "heading",
    [
        "文献报告准确率为99.9%",
        "1. 文献报告准确率为99.9%",
        "1.2. 文献报告准确率为99.9%",
        "第2章 文献报告准确率为99.9%",
        "（3）文献报告准确率为99.9%",
    ],
)
def test_cited_numbers_and_structural_numbering_are_allowed_in_headings(
    store, kind, heading
):
    sid = abstract_source(store, "The cited study reports accuracy of 99.9%.")
    parent = write_draft(store, sid) if kind == "revision" else None

    def result(payload):
        data = writing_result(payload)
        data["sections"][0]["heading"] = heading
        return data

    kwargs = {"parent_id": parent["id"]} if parent else {}
    draft = write_draft(store, sid, Backend(result), kind=kind, **kwargs)
    assert heading in draft["markdown"]
    assert draft["references"][0]["source_id"] == sid
    assert draft["references"][0]["bibliography_status"] == "unverified"
    assert store.rows("SELECT * FROM evidence_cards") == []
    run = store.rows("SELECT * FROM model_runs WHERE id=?", (draft["model_run_id"],))[0]
    assert run["status"] == "success"


@pytest.mark.parametrize(
    "heading",
    [
        "1. 研究背景",
        "1.2. 研究背景",
        "2、研究背景",
        "第3章 研究背景",
        "（4）研究背景",
    ],
)
def test_pure_chapter_numbers_do_not_require_numeric_evidence(store, heading):
    sid = abstract_source(store)

    def result(payload):
        data = writing_result(payload)
        data["sections"][0]["heading"] = heading
        return data

    draft = write_draft(store, sid, Backend(result))
    assert heading in draft["markdown"]
    assert paper_work.get_state(store)["writing_drafts"] == [draft]


@pytest.mark.parametrize("target", ["title", "heading"])
def test_title_uses_whole_draft_citations_but_heading_only_uses_its_section(
    store, target
):
    sid = abstract_source(
        store, "The cited study reports accuracy of 99.9%. Context limits transfer."
    )

    def result(payload):
        data = writing_result(payload)
        citation_base = citation(payload)
        data["sections"][0]["paragraphs"][0]["citations"] = [
            {**citation_base, "quote": "Context limits transfer."}
        ]
        data["sections"].append(
            {
                "heading": "文献结果",
                "paragraphs": [
                    {
                        "text": "文献报告准确率为99.9%。",
                        "citations": [
                            {
                                **citation_base,
                                "quote": "The cited study reports accuracy of 99.9%.",
                            }
                        ],
                    }
                ],
            }
        )
        (data if target == "title" else data["sections"][0])[target] = (
            "文献报告准确率为99.9%"
        )
        return data

    if target == "title":
        draft = write_draft(store, sid, Backend(result))
        assert draft["title"] == "文献报告准确率为99.9%"
    else:
        with pytest.raises(ValueError):
            write_draft(store, sid, Backend(result))
        assert_failed(store)
        assert paper_work.get_state(store)["writing_drafts"] == []


@pytest.mark.parametrize("target", ["title", "heading"])
def test_heading_numbers_cannot_borrow_uncited_input_text(store, target):
    sid = abstract_source(
        store, "The cited study reports accuracy of 99.9%. Context limits transfer."
    )

    def result(payload):
        data = writing_result(payload)
        data["sections"][0]["paragraphs"][0]["citations"][0]["quote"] = (
            "Context limits transfer."
        )
        (data if target == "title" else data["sections"][0])[target] = (
            "文献报告准确率为99.9%"
        )
        return data

    with pytest.raises(ValueError):
        write_draft(store, sid, Backend(result))
    assert_failed(store)
    assert paper_work.get_state(store)["writing_drafts"] == []


def test_numeric_outline_plan_remains_structure_without_evidence(store):
    backend = Backend(
        {
            "title": "3步研究计划",
            "sections": [
                {
                    "heading": "1. 拟采样100例",
                    "paragraphs": [
                        {"text": "计划抽样100例，尚待可行性论证。", "citations": []}
                    ],
                }
            ],
        }
    )
    draft = paper_work.generate_writing(
        store,
        backend,
        kind="outline",
        instructions="拟定研究结构",
        source_ids=[],
        rights_confirmed=False,
    )
    assert draft["kind"] == "outline"
    assert draft["references"] == [] and draft["coverage"] == []
    assert store.rows("SELECT * FROM evidence_cards") == []
    refs = json.loads(store.rows("SELECT * FROM model_runs")[0]["input_refs"])
    assert refs["source_ids"] == [] and refs["pages"] == []
    assert refs["excerpts"] == [] and refs["source_fingerprints"] == {}


def test_revision_uses_requested_parent_and_preserves_ai_history(store):
    sid = abstract_source(store)
    first = write_draft(store, sid)
    backend = Backend(writing_result)
    revised = write_draft(store, sid, backend, kind="revision", parent_id=first["id"])
    assert revised["parent_id"] == first["id"] and revised["kind"] == "revision"
    assert backend.calls[0]["payload"]["parent_markdown"] == first["markdown"]
    drafts = paper_work.get_state(Store(store.data_dir))["writing_drafts"]
    assert drafts == [first, revised]


def test_revision_reduces_excerpt_budget_instead_of_truncating_parent(store):
    sid = abstract_source(store, "learning " * 6000)
    first = write_draft(store, sid)
    backend = Backend(writing_result)
    revised = write_draft(store, sid, backend, kind="revision", parent_id=first["id"])
    payload = backend.calls[0]["payload"]
    assert payload["parent_markdown"] == first["markdown"]
    assert (
        sum(len(item["text"]) for item in payload["excerpts"])
        + len(payload["parent_markdown"])
        <= 24000
    )
    assert revised["coverage"][0]["characters"] < first["coverage"][0]["characters"]


def test_manual_edit_creates_version_and_never_overwrites_model_or_duplicates_references(
    store,
):
    sid = abstract_source(store)
    first = write_draft(store, sid)
    run = store.rows("SELECT * FROM model_runs")[0]
    edited = paper_work.save_writing(
        store,
        first["id"],
        markdown=first["markdown"].replace("文献讨论了", "人工修改：文献讨论了"),
        title="人工修订标题",
    )
    assert edited["parent_id"] == first["id"] and edited["kind"] == "manual"
    assert edited["model_run_id"] is None
    assert edited["references"] == first["references"]
    assert edited["coverage"] == first["coverage"]
    assert edited["markdown"].count("## 参考文献") == 1
    assert "人工修改" in edited["markdown"] and "人工修改" not in first["markdown"]
    assert store.rows("SELECT * FROM model_runs") == [run]
    assert paper_work.get_state(Store(store.data_dir))["writing_drafts"] == [
        first,
        edited,
    ]


def test_manual_edits_cannot_forge_reference_suffix_or_unknown_number(store):
    sid = abstract_source(store)
    first = write_draft(store, sid)
    for content in [
        first["markdown"].replace("Known Author", "Fake Author"),
        "伪引用[999]",
        "## 参考文献\n伪造条目",
    ]:
        with pytest.raises(ValueError):
            paper_work.save_writing(store, first["id"], markdown=content)
    assert paper_work.get_state(store)["writing_drafts"] == [first]


def test_historical_readings_and_drafts_survive_but_cannot_be_saved_as_current(store):
    sid = abstract_source(store)
    report = read(store, sid)
    draft = write_draft(store, sid)
    confirm(store, "另一个研究问题")
    with pytest.raises(paper_work.PaperWorkConflict):
        paper_work.save_writing(store, draft["id"], markdown="新的正文[1]")
    backend = Backend(writing_result)
    with pytest.raises(paper_work.PaperWorkConflict):
        write_draft(store, sid, backend, kind="revision", parent_id=draft["id"])
    assert backend.calls == []
    assert paper_work.get_state(Store(store.data_dir)) == {
        "reading_reports": [report],
        "writing_drafts": [draft],
    }
    for fmt in ["md", "bib"]:
        assert (
            "历史研究版本，不是当前成果"
            in paper_work.export_writing(store, draft["id"], fmt).decode()
        )
    doc = Document(BytesIO(paper_work.export_writing(store, draft["id"], "docx")))
    assert any(
        "历史研究版本，不是当前成果" in paragraph.text for paragraph in doc.paragraphs
    )


def test_real_markdown_word_and_bibtex_exports_include_content_and_snapshot(store):
    sid = abstract_source(store)
    draft = write_draft(store, sid)
    md = paper_work.export_writing(store, draft["id"], "md").decode()
    assert "# 学习环境研究草稿" in md
    assert draft["markdown"] in md
    doc = Document(BytesIO(paper_work.export_writing(store, draft["id"], "docx")))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "学习环境研究草稿" in text and "文献讨论了学习环境及其适用边界。[1]" in text
    assert "Known Author" in text and "AI辅助草稿" in text and "仅摘要" in text
    assert (
        doc.core_properties.author == "" and doc.core_properties.last_modified_by == ""
    )
    bib = paper_work.export_writing(store, draft["id"], "bib").decode()
    assert (
        "@misc{ref1," in bib
        and "author = {Known Author}" in bib
        and "doi = {10.1234/known}" in bib
    )
    # 来源后来变更，导出仍保留生成当时书目，不悄悄改成新引用。
    with store.connection() as db:
        db.execute(
            "UPDATE sources SET title='Changed after generation' WHERE id=?", (sid,)
        )
    exported = paper_work.export_writing(store, draft["id"], "md").decode()
    assert (
        "Learning evidence" in exported and "Changed after generation" not in exported
    )
    assert "来源完整性已变化" in exported


def test_untrusted_metadata_and_manual_body_are_inert_in_exports(store):
    sid = abstract_source(
        store, title='<script>alert("x")</script> {fake} \\input{evil}'
    )
    draft = write_draft(store, sid)
    assert "<script>" not in draft["markdown"]
    manual = paper_work.save_writing(
        store,
        draft["id"],
        markdown="<script>alert(1)</script>\n[bad](javascript:alert(1))\n![image](https://tracking.invalid)\n真实引文[1]",
    )
    assert "<script>" not in manual["markdown"]
    assert "\\[bad]" in manual["markdown"] and "!\\[image]" in manual["markdown"]
    bib = paper_work.export_writing(store, draft["id"], "bib").decode()
    assert r"\{fake\}" in bib and r"\textbackslash{}input\{evil\}" in bib
    document = Document(BytesIO(paper_work.export_writing(store, manual["id"], "docx")))
    assert all(
        "javascript:" not in rel.target_ref for rel in document.part.rels.values()
    )


def test_invalid_export_format_and_missing_versions_fail(store):
    sid = abstract_source(store)
    draft = write_draft(store, sid)
    with pytest.raises(ValueError):
        paper_work.export_writing(store, draft["id"], "../md")
    with pytest.raises(ValueError):
        paper_work.export_writing(store, "missing", "md")


def test_duplicate_sources_are_sent_once_and_not_counted_as_independent_support(store):
    sid = abstract_source(store)
    backend = Backend()
    report = paper_work.read_sources(
        store,
        backend,
        source_ids=[sid, sid],
        question="learning",
        rights_confirmed=True,
    )
    assert len(report["coverage"]) == 1
    assert len(backend.calls[0]["payload"]["excerpts"]) == 1
    writing = Backend(writing_result)
    draft = paper_work.generate_writing(
        store,
        writing,
        kind="draft",
        instructions="背景草稿",
        source_ids=[sid, sid],
        rights_confirmed=True,
    )
    assert len(draft["references"]) == 1 and len(draft["coverage"]) == 1


def test_numeric_citations_cannot_be_converted_to_active_links(store):
    sid = abstract_source(store)
    draft = write_draft(store, sid)
    manual = paper_work.save_writing(
        store,
        draft["id"],
        markdown="[1](javascript:alert(1))\n[1]: https://tracking.invalid\n保留引用[1]",
    )
    assert r"\[1](javascript:" in manual["markdown"]
    assert r"\[1]: https://" in manual["markdown"]
    assert "保留引用[1]" in manual["markdown"]
