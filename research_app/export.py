"""匿名过程导出与参赛文件门禁；技术检查不替代人的原文、版面及逐帧复核。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import tempfile
import unicodedata
from io import BytesIO, StringIO
from pathlib import Path
from xml.sax.saxutils import escape

from . import (
    decision,
    discovery,
    evaluation,
    evidence,
    experiment,
    paper_work,
    web_research,
)
from .models import now
from .storage import Store

TITLE = "从“AI说”到“证据说”：科研选题的可追溯论证实践"
PREVIEW = "排版预览，非提交版"
PREVIEW_SONG = ("/System/Library/Fonts/Supplemental/Songti.ttc", 6)
PREVIEW_HEI = ("/System/Library/Fonts/STHeiti Medium.ttc", 1)
BRIEF_FIELDS = (
    "question",
    "population_or_object",
    "constraints",
    "available_resources",
    "validation_goal",
)
EVIDENCE_FIELDS = (
    "claim",
    "source_id",
    "page_number",
    "quote",
    "conditions",
    "limitations",
    "ai_relation",
)
SOURCE_FIELDS = (
    "id",
    "title",
    "authors",
    "year",
    "doi",
    "official_url",
    "access_level",
    "bibliography_status",
    "sha256",
    "page_count",
    "inclusion_reason",
    "exclusion_reason",
    "created_at",
)
REVIEW_FIELDS = (
    "id",
    "method",
    "question_id",
    "claim_index",
    "reviewer_id",
    "source_locatable",
    "source_supports",
    "within_scope",
    "correction_count",
    "note",
    "review_seconds",
    "reviewed_at",
)
SENSITIVE = re.compile(
    r"(?:/(?:Users|home|private|tmp|var|opt|Volumes|etc)/|\b[A-Za-z]:[\\/]|file://|\bBearer\s+\S+|\b(?:sk|sess)-[A-Za-z0-9_-]{8,}|\b(?:api[_ -]?key|access[_ -]?token|refresh[_ -]?token|password|authorization)\s*[\"']?\s*[:=]\s*\S+|https?://[^/\s]+@|[\w.+-]+@[\w.-]+\.[A-Za-z]{2,})",
    re.I,
)
IDENTITY_LABEL = re.compile(
    r"(?:学校|指导教师|学生姓名|团队名称|参赛者姓名)\s*[:：]\s*\S+"
)


class ExportError(ValueError):
    """安全的导出错误，不包含本机路径或底层异常。"""


def _pick(value: dict, fields) -> dict:
    return {key: value[key] for key in fields if key in value}


def _clean(value, redactions: list[str], field="", sms=()):
    if isinstance(value, str):
        value = unicodedata.normalize("NFC", value)
        if (
            SENSITIVE.search(value)
            or IDENTITY_LABEL.search(value)
            or any(
                message == value.strip() or (len(message) >= 20 and message in value)
                for message in sms
            )
        ):
            redactions.append(field)
            return "[敏感字段已移除，须人工核对]"
        return "".join(
            c for c in value if c in "\n\t" or unicodedata.category(c) != "Cc"
        )
    if isinstance(value, dict):
        return {k: _clean(v, redactions, f"{field}.{k}", sms) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [
            _clean(v, redactions, f"{field}[{i}]", sms) for i, v in enumerate(value)
        ]
    return value


def _bytes(value) -> bytes:
    return (
        json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False)
        + "\n"
    ).encode("utf-8")


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _safe_path(path: str | Path) -> Path:
    path = Path(path).absolute()
    if any(p.is_symlink() for p in (path, *path.parents)):
        raise ExportError("输出路径及其父目录不允许符号链接")
    return path


def _directory(path: str | Path, *, draft=False) -> Path:
    path = _safe_path(path)
    if draft and "submission" in path.parts:
        raise ExportError("过程记录和预览不能写入正式submission目录")
    path.mkdir(parents=True, exist_ok=True)
    return path


def _write(path: Path, data: bytes, *, no_replace=False) -> None:
    path = _safe_path(path)
    if no_replace and path.exists():
        if path.is_file() and path.read_bytes() == data:
            return
        raise ExportError("目标已存在且内容不同，拒绝覆盖")
    if no_replace:
        try:
            with path.open("xb") as handle:
                handle.write(data)
        except FileExistsError:
            raise ExportError("目标已存在，拒绝覆盖") from None
        return
    fd, temporary = tempfile.mkstemp(prefix=".export-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def font_preflight(path: str | Path, index=0, *, expected: str | None = None) -> dict:
    """检查真实字体名称及嵌入许可；只读字体源，不伪造字体名。"""
    from fontTools.ttLib import TTFont

    if type(index) is not int or index < 0:
        raise ExportError("字体集合索引必须为非负整数")
    path = Path(path)
    if path.suffix.lower() not in {".ttf", ".ttc"} or not path.is_file():
        raise ExportError("请提供合法的TTF或TTC字体文件")
    if path.suffix.lower() == ".ttf" and index != 0:
        raise ExportError("单字体TTF的索引必须为0")
    try:
        with TTFont(str(path), fontNumber=index, lazy=False) as font:
            names = {
                str(key): sorted(
                    {
                        record.toUnicode()
                        for record in font["name"].names
                        if record.nameID == key
                    }
                )
                for key in (1, 2, 4, 6)
            }
            fs_type = int(font["OS/2"].fsType)
            outline = "glyf" in font
    except Exception:
        raise ExportError("无法读取指定字体、名称表或集合索引") from None
    allowed = outline and not (fs_type & (0x2 | 0x100 | 0x200))
    authentic = expected is None or expected in names["6"]
    reasons = []
    if not allowed:
        reasons.append("字体禁止所需嵌入/子集化，或并非支持的TrueType轮廓")
    if not authentic:
        reasons.append(f"真实PostScript名称不是{expected}，不能冒充比赛指定字体")
    return {
        "names": names,
        "fs_type": fs_type,
        "collection_index": index,
        "true_type": outline,
        "embeddable": allowed,
        "matches_expected": authentic,
        "ok": allowed and authentic,
        "reasons": reasons,
    }


def _fonts(song, hei, song_index, hei_index):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    result = []
    for path, index in ((song, song_index), (hei, hei_index)):
        info = font_preflight(path, index)
        if not info["ok"]:
            raise ExportError("所选字体不允许嵌入或不支持PDF输出")
        # 路径只用于本机注册键，绝不进入PDF元数据或导出记录。
        name = "Export" + _digest(f"{Path(path).absolute()}:{index}".encode())[:16]
        if name not in pdfmetrics.getRegisteredFontNames():
            try:
                pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=index))
            except Exception:
                raise ExportError("字体无法用于PDF排版") from None
        result.append(name)
    return result


def _pdf(
    title: str, sections: list[tuple[str, str]], *, preview=True, fonts=None
) -> bytes:
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    if not fonts:
        body_font = os.environ.get("RESEARCH_PREVIEW_BODY_FONT")
        heading_font = os.environ.get("RESEARCH_PREVIEW_HEADING_FONT")
        fonts = _fonts(
            body_font or PREVIEW_SONG[0],
            heading_font or PREVIEW_HEI[0],
            0 if body_font else PREVIEW_SONG[1],
            0 if heading_font else PREVIEW_HEI[1],
        )
    song, hei = fonts
    output = BytesIO()
    body = ParagraphStyle(
        "Body",
        fontName=song,
        fontSize=14,
        leading=21,
        firstLineIndent=28,
        wordWrap="CJK",
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "Title",
        fontName=hei,
        fontSize=16,
        leading=24,
        alignment=TA_CENTER,
        wordWrap="CJK",
        spaceAfter=12,
    )
    label = ParagraphStyle(
        "Label", fontName=hei, fontSize=14, leading=21, wordWrap="CJK", spaceAfter=8
    )
    story = [Paragraph(escape(title), heading)]
    if preview:
        story.append(Paragraph(PREVIEW, label))
    for name, text in sections:
        if name:
            story.append(Paragraph(escape(name), label))
        for paragraph in text.split("\n"):
            if paragraph.strip():
                story.append(Paragraph(escape(paragraph), body))
        story.append(Spacer(1, 5))

    def canvasmaker(*args, **kwargs):
        kwargs["initialFontName"] = song
        kwargs["invariant"] = 1
        canvas = Canvas(*args, **kwargs)
        canvas.setAuthor("")
        canvas.setCreator("")
        canvas.setSubject("")
        canvas.setTitle(title)
        return canvas

    SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=25 * mm,
        leftMargin=25 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm,
        title=title,
        author="",
        creator="",
        subject="",
    ).build(story, canvasmaker=canvasmaker)
    return output.getvalue()


def _aggregate(rows: list[dict]) -> list[dict]:
    result = []
    for row in rows:
        item = _pick(
            row, ("candidate", "budget", "seeds", "spam_support", "ham_support")
        )
        for group, allowed in (
            ("metrics", experiment.METRICS),
            ("timing", experiment.TIMINGS),
        ):
            item[group] = {
                metric: _pick(
                    row.get(group, {}).get(metric, {}), ("mean", "sample_std", "n")
                )
                for metric in allowed
                if metric in row.get(group, {})
            }
        result.append(item)
    return result


def _experiments(store: Store) -> list[dict]:
    result = []
    _safe_path(store.data_dir / "experiments")
    for row in experiment.list_runs(store.data_dir):
        if not row["verified"]:
            result.append(
                {
                    "run_id": row["run_id"],
                    "status": "unverified",
                    "note": "完整性校验未通过，不导出指标",
                }
            )
            continue
        try:
            run_path = _safe_path(store.data_dir / "experiments" / row["run_id"])
            for name in (
                *experiment.BASE_FILES,
                *experiment.OPTIONAL_FILES,
                "integrity.json",
            ):
                _safe_path(run_path / name)
            summary = experiment.verified_summary(
                row["run_id"], data_dir=store.data_dir
            )
        except ValueError:
            result.append({"run_id": row["run_id"], "status": "unverified"})
            continue
        item = _pick(
            summary,
            (
                "run_id",
                "created_at",
                "status",
                "snapshot_sha256",
                "config_sha256",
                "split_sha256",
                "recommended_candidate",
                "validation_completed_at",
                "test_completed_at",
                "validation_fit_count",
                "test_fit_count",
            ),
        )
        for group in ("validation_aggregates", "test_aggregates"):
            item[group] = _aggregate(summary.get(group, []))
        item["actual_timing"] = _pick(
            summary.get("actual_timing", {}),
            (
                "validation_wall_seconds",
                "validation_fit_seconds",
                "test_wall_seconds",
                "test_fit_seconds",
            ),
        )
        item["config"] = experiment._config()
        manifest = json.loads((run_path / "manifest.json").read_text(encoding="utf-8"))
        item["data_audit"] = _pick(
            manifest,
            (
                "audited_at",
                "source_url",
                "license",
                "official_sample_count",
                "snapshot_sha256",
                "status",
                "actual_parsed_count",
                "difference_from_official",
                "retained_count",
                "duplicate_removed_count",
                "conflict_group_count",
                "conflict_row_count",
                "count_note",
                "limitation",
            ),
        )
        item["data_audit"]["label_counts"] = _pick(
            manifest.get("label_counts", {}), ("ham", "spam")
        )
        item["data_audit"]["retained_label_counts"] = _pick(
            manifest.get("retained_label_counts", {}), ("ham", "spam")
        )
        item["paired_differences"] = []
        for pair in summary.get("paired_differences", []):
            item["paired_differences"].append(
                {
                    **_pick(pair, ("budget", "direction")),
                    "metrics": {
                        key: {
                            **_pick(value, ("mean", "sample_std", "n")),
                            "by_seed": _pick(
                                value.get("by_seed", {}), ("0", "1", "2", "3", "4")
                            ),
                        }
                        for key, value in pair.get("metrics", {}).items()
                        if key in experiment.METRICS
                    },
                }
            )
        if "frozen_selection" in summary:
            item["frozen_selection"] = _pick(
                summary["frozen_selection"],
                (
                    "run_id",
                    "chosen_candidate",
                    "study_revision",
                    "reason",
                    "confirmed_at",
                    "snapshot_sha256",
                    "config_sha256",
                    "split_sha256",
                ),
            )
        result.append(item)
    return result


def process_snapshot(
    store: Store, *, source_checks: dict[str, bool] | None = None
) -> dict:
    """白名单快照；source_checks仅可在当前请求内与卡片资格核验共享。"""
    brief = store.latest_brief()
    revision = brief["revision"] if brief else None
    briefs = []
    for row in store.rows(
        "SELECT id,revision,content,queries,confirmed_at FROM briefs ORDER BY revision"
    ):
        briefs.append(
            {
                **_pick(row, ("id", "revision", "confirmed_at")),
                "content": _pick(json.loads(row["content"]), BRIEF_FIELDS),
                "queries": json.loads(row["queries"]),
            }
        )
    cards = []
    if source_checks is None:
        source_checks = {}
    for card in evidence.list_cards(store):
        with store.connection() as db:
            matched, note = evidence._validate_quote(
                db, card["content"], card["analysis_id"]
            )
        cards.append(
            {
                **_pick(
                    card,
                    (
                        "id",
                        "analysis_id",
                        "study_revision",
                        "review_status",
                        "review_note",
                        "reviewed_at",
                        "created_at",
                    ),
                ),
                "content": _pick(card["content"], EVIDENCE_FIELDS),
                "quote_match": matched,
                "validation_note": note,
                "eligible": evidence.eligible_card(
                    store, card["id"], revision=revision, source_checks=source_checks
                ),
            }
        )
    card_by_id = {card["id"]: card for card in cards}
    candidates = []
    for row in decision.list_candidate_runs(store):
        content = row
        candidates.append(
            {
                **_pick(row, ("id", "study_revision", "created_at")),
                "sms_case": content.get("sms_case"),
                "input_evidence_ids": content.get("input_evidence_ids", []),
                "candidates": [
                    _pick(
                        c,
                        (
                            "candidate_id",
                            "method",
                            "evidence_ids",
                            "conditions",
                            "feasibility",
                            "minimal_validation",
                            "objections",
                        ),
                    )
                    for c in content.get("candidates", [])
                ],
            }
        )
    candidate_by_id = {row["id"]: row for row in candidates}
    experiments = _experiments(store)
    runs = {row["run_id"]: row for row in experiments}
    decisions = []
    for row in decision.list_decisions(store):
        reasons = []
        if row["status"] != "final" or row.get("invalidated_at"):
            reasons.append("记录为草案或已失效")
        if revision is None or row["study_revision"] != revision:
            reasons.append("非当前已确认研究修订")
        if not row["evidence_ids"] or not all(
            card_by_id.get(i, {}).get("eligible") for i in row["evidence_ids"]
        ):
            reasons.append("存在未通过当前书目、原PDF完整性、原句和人工支持核验的证据")
        candidate = candidate_by_id.get(row.get("candidate_run_id"))
        chosen = (
            next(
                (
                    c
                    for c in candidate["candidates"]
                    if c["candidate_id"] == row["chosen_candidate"]
                ),
                None,
            )
            if candidate
            else None
        )
        if (
            not candidate
            or candidate["study_revision"] != revision
            or not chosen
            or not set(row["evidence_ids"]).issubset(chosen["evidence_ids"])
        ):
            reasons.append("候选来源、研究修订或引用关联不成立")
        if not row["reason"].strip():
            reasons.append("缺少人工选择理由")
        experiment_reason = decision.experiment_gate_reason(
            brief,
            candidate,
            row["chosen_candidate"],
            row.get("experiment_run_id"),
            runs.get(row.get("experiment_run_id")),
        )
        if experiment_reason:
            reasons.append(experiment_reason)
        decisions.append(
            {
                **_pick(
                    row,
                    (
                        "id",
                        "study_revision",
                        "chosen_candidate",
                        "reason",
                        "evidence_ids",
                        "experiment_run_id",
                        "status",
                        "created_at",
                        "invalidated_at",
                        "candidate_run_id",
                    ),
                ),
                "effective_status": "final" if not reasons else "draft",
                "gate_reasons": reasons,
            }
        )
    reviews = []
    for row in store.rows(
        "SELECT id,card_id,previous_content,previous_status,review_note,created_at FROM evidence_reviews ORDER BY created_at,id"
    ):
        reviews.append(
            {
                **_pick(
                    row,
                    ("id", "card_id", "previous_status", "review_note", "created_at"),
                ),
                "previous_content": _pick(
                    json.loads(row["previous_content"]), EVIDENCE_FIELDS
                ),
            }
        )
    model_runs = []
    for row in store.rows(
        "SELECT id,study_revision,task,backend,requested_model,actual_model,called_at,prompt_version,input_refs,status,error_type FROM model_runs ORDER BY called_at,id"
    ):
        item = _pick(
            row,
            (
                "id",
                "study_revision",
                "task",
                "backend",
                "requested_model",
                "actual_model",
                "called_at",
                "prompt_version",
                "status",
                "error_type",
            ),
        )
        refs = json.loads(row["input_refs"])
        item["input_refs"] = _pick(
            refs,
            (
                "source_ids",
                "evidence_ids",
                "candidate_run_id",
                "experiment_run_id",
                "search_run_id",
                "run_id",
                "plan_id",
                "parent_id",
                "draft_id",
                "web_research_id",
                "kind",
            ),
        )
        item["input_refs"]["pages"] = [
            _pick(p, ("source_id", "page_number")) for p in refs.get("pages", [])
        ]
        if "coverage" in refs:
            item["input_refs"]["coverage"] = [
                _pick(
                    covered,
                    (
                        "source_id",
                        "title",
                        "level",
                        "pages",
                        "total_pages",
                        "characters",
                    ),
                )
                for covered in refs["coverage"]
            ]
        if "source_fingerprints" in refs:
            item["input_refs"]["source_fingerprints"] = {
                source_id: digest
                for source_id, digest in refs["source_fingerprints"].items()
                if isinstance(digest, str) and re.fullmatch(r"[0-9a-f]{64}", digest)
            }
        if "excerpts" in refs:
            item["input_refs"]["excerpts"] = [
                _pick(
                    excerpt,
                    (
                        "source_id",
                        "page_number",
                        "level",
                        "start_character",
                        "end_character",
                        "text_sha256",
                    ),
                )
                for excerpt in refs["excerpts"]
            ]
        if "web_pages" in refs:
            item["input_refs"]["web_pages"] = [
                _pick(
                    page,
                    (
                        "id",
                        "url",
                        "final_url",
                        "retrieved_at",
                        "characters",
                        "excerpt_sha256",
                        "content_sha256",
                    ),
                )
                for page in refs["web_pages"]
            ]
        model_runs.append(item)
    evaluations = []
    for record in evaluation.list_evaluations(store):
        item = _pick(
            record,
            (
                "id",
                "created_at",
                "study_revision",
                "experiment_run_id",
                "status",
                "current_revision",
                "pending_claim_count",
                "packet_sha256",
                "scope_note",
            ),
        )
        item["methods"] = {
            method: _pick(
                state,
                (
                    "status",
                    "requested_model",
                    "actual_model",
                    "model_seconds",
                    "review_seconds",
                    "total_seconds",
                ),
            )
            for method, state in record["methods"].items()
            if method in evaluation.METHODS
        }
        item["reviews"] = [_pick(review, REVIEW_FIELDS) for review in record["reviews"]]
        evaluations.append(item)
    search = []
    for row in store.rows(
        "SELECT id,study_revision,query,requested_at,status,error_type,screening_json FROM search_runs ORDER BY requested_at,id"
    ):
        item = _pick(
            row,
            ("id", "study_revision", "query", "requested_at", "status", "error_type"),
        )
        item["screening"] = [
            _pick(
                result,
                (
                    "title",
                    "authors",
                    "year",
                    "doi",
                    "official_url",
                    "access_level",
                    "include",
                    "reason",
                    "source_id",
                ),
            )
            for result in json.loads(row["screening_json"])
        ]
        search.append(item)
    discovery_state = discovery.get_state(store)
    literature_searches = []
    for run in discovery_state["runs"]:
        item = _pick(
            run,
            (
                "id",
                "study_revision",
                "provider",
                "query",
                "status",
                "message",
                "created_at",
            ),
        )
        item["results"] = [
            _pick(
                result,
                (
                    "index",
                    "title",
                    "authors",
                    "year",
                    "doi",
                    "official_url",
                    "remote_id",
                    "access_level",
                    "source_id",
                    "recommendation",
                ),
            )
            for result in run["results"]
        ]
        literature_searches.append(item)
    writing_state = paper_work.get_state(store)
    reading_records = [
        _pick(
            report,
            (
                "id",
                "study_revision",
                "question",
                "coverage",
                "model_run_id",
                "created_at",
            ),
        )
        for report in writing_state["reading_reports"]
    ]
    writing_records = []
    for draft in writing_state["writing_drafts"]:
        item = _pick(
            draft,
            (
                "id",
                "study_revision",
                "kind",
                "title",
                "references",
                "coverage",
                "model_run_id",
                "parent_id",
                "created_at",
            ),
        )
        item["content_sha256"] = hashlib.sha256(draft["markdown"].encode()).hexdigest()
        item["notice"] = (
            "写作版本记录，不表示引用已核验或论文已完成；正文另由草稿导出。"
        )
        writing_records.append(item)
    web_records = []
    for run in web_research.get_state(store)["runs"]:
        item = _pick(
            run,
            (
                "id",
                "study_revision",
                "question",
                "status",
                "created_at",
                "completed_at",
                "provider",
                "message",
                "model_run_ids",
            ),
        )
        item["queries"] = [
            _pick(query, ("query", "status", "message", "searched_at"))
            | {
                "results": [
                    _pick(result, ("url", "title")) for result in query["results"]
                ]
            }
            for query in run["queries"]
        ]
        answer_call = next(
            (
                call
                for call in model_runs
                if call["task"] == "web_research_answer"
                and call["input_refs"].get("web_research_id") == run["id"]
            ),
            None,
        )
        input_pages = {
            page["id"]: page["characters"]
            for page in (
                answer_call["input_refs"].get("web_pages", []) if answer_call else []
            )
        }
        item["answer_model_run_id"] = answer_call["id"] if answer_call else None
        item["pages"] = [
            _pick(
                page,
                (
                    "id",
                    "url",
                    "final_url",
                    "title",
                    "status",
                    "message",
                    "retrieved_at",
                    "extracted_characters",
                    "truncated",
                    "content_sha256",
                    "excerpt_sha256",
                ),
            )
            | {
                "prepared_characters": len(page["excerpt"]),
                "model_input_characters": input_pages.get(page["id"], 0),
            }
            for page in run["pages"]
        ]
        item["notice"] = (
            "联网查阅过程，不含网页正文或AI原始回答；待用片段不代表已发给模型，模型输入清单也不保证失败调用已送达服务端。来源与结论仍待人工核验。"
        )
        web_records.append(item)
    snapshot = {
        "schema_version": 1,
        "current_revision": revision,
        "status": "formal_decision_available"
        if any(d["effective_status"] == "final" for d in decisions)
        else "incomplete",
        "notice": "研究过程记录，不是参赛说明。引文命中不等于推论获支持；未人工核验的内容不是科研结论。机器筛查不等于人工匿名确认。",
        "briefs": briefs,
        "search_runs": search,
        "literature_plans": discovery_state["plans"],
        "literature_searches": literature_searches,
        "reading_records": reading_records,
        "writing_records": writing_records,
        "web_research_records": web_records,
        "sources": [_pick(s, SOURCE_FIELDS) for s in store.list_sources()],
        "evidence": cards,
        "human_review_history": reviews,
        "objections": [c["id"] for c in cards if not c["eligible"]],
        "candidate_runs": candidates,
        "decisions": decisions,
        "experiments": experiments,
        "model_runs": model_runs,
        "evaluations": evaluations,
    }
    redactions = []
    sms = ()
    corpus = _safe_path(store.data_dir / "case" / "SMSSpamCollection")
    if corpus.is_file():
        try:
            sms = tuple(
                line.split("\t", 1)[1].strip()
                for line in corpus.read_text(encoding="utf-8").splitlines()
                if "\t" in line and line.split("\t", 1)[1].strip()
            )
        except (OSError, UnicodeError):
            raise ExportError("短信快照无法读取，不能完成正文泄露筛查") from None
    snapshot = _clean(snapshot, redactions, sms=sms)
    snapshot["privacy"] = {
        "human_anonymity_review_required": True,
        "redacted_fields": redactions,
        "excluded": [
            "模型原始正文",
            "PDF全文",
            "网页正文与联网回答原文",
            "短信正文",
            "API密钥及账户配置",
            "本机路径",
        ],
    }
    return snapshot


def _csv_cell(value):
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(
        ("\t", "\r", "\n")
    ):
        return "'" + text
    return text


def _readable(value, prefix="") -> list[str]:
    lines = []
    if isinstance(value, dict):
        for key, item in value.items():
            lines.extend(_readable(item, f"{prefix}{key}："))
    elif isinstance(value, list):
        if not value:
            lines.append(prefix + "无记录")
        for index, item in enumerate(value, 1):
            lines.extend(_readable(item, f"{prefix}{index}. "))
    else:
        lines.append(prefix + ("未记录" if value is None else str(value)))
    return lines


def export_process(store: Store, output_dir="artifacts/process") -> dict:
    root = _directory(output_dir, draft=True)
    snapshot = process_snapshot(store)
    _write(root / "process.json", _bytes(snapshot))
    fields = (
        "id",
        "study_revision",
        *EVIDENCE_FIELDS,
        "quote_match",
        "review_status",
        "review_note",
        "reviewed_at",
        "eligible",
        "validation_note",
    )
    buffer = StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=fields)
    writer.writeheader()
    for card in snapshot["evidence"]:
        flat = {**card, **card["content"]}
        writer.writerow({field: _csv_cell(flat.get(field)) for field in fields})
    _write(root / "evidence.csv", buffer.getvalue().encode("utf-8-sig"))
    sections = [("记录状态", snapshot["status"] + "\n" + snapshot["notice"])]
    for key, label in (
        ("briefs", "研究问题与修订"),
        ("search_runs", "真实检索审计"),
        ("literature_plans", "AI检索策略"),
        ("literature_searches", "公开文献库真实查询与纳入"),
        ("web_research_records", "公开网页联网查阅、覆盖与失败审计"),
        ("reading_records", "实际阅读覆盖与调用关联"),
        ("writing_records", "写作版本、引用与人工修改链"),
        ("sources", "来源与访问层级"),
        ("evidence", "证据、原句与人工判断"),
        ("human_review_history", "人工修改与复核日志"),
        ("candidate_runs", "候选比较与异议"),
        ("decisions", "决策及动态资格"),
        ("experiments", "已核验实验与哈希"),
        ("model_runs", "模型身份与调用审计"),
        ("evaluations", "评价进度与人工记录"),
        ("privacy", "导出边界"),
    ):
        sections.append((label, "\n".join(_readable(snapshot[key]))))
    files = ["process.json", "evidence.csv"]
    blocked = []
    try:
        _write(
            root / "process.pdf", _pdf("研证：研究过程记录", sections, preview=False)
        )
        files.append("process.pdf")
    except ExportError as exc:
        blocked.append(str(exc))
    return {
        "status": snapshot["status"],
        "files": files,
        "blocked": blocked,
        "formal_decision_count": sum(
            d["effective_status"] == "final" for d in snapshot["decisions"]
        ),
        "human_anonymity_review_required": True,
    }


def character_count(text: str) -> int:
    """保守按所有非空白Unicode字符计数，数字、英文、标点和声明均计入。"""
    return len(re.sub(r"\s", "", text))


def _case_content(snapshot: dict) -> tuple[list[str], list[str], list[str]]:
    current = snapshot["current_revision"]
    formal = [d for d in snapshot["decisions"] if d["effective_status"] == "final"]
    confirmed = [
        r
        for r in snapshot["experiments"]
        if r.get("status") == "confirmed"
        and r.get("frozen_selection", {}).get("study_revision") == current
    ]
    linked = {d.get("experiment_run_id") for d in formal}
    run = next(
        (r for r in confirmed if r["run_id"] in linked),
        confirmed[0] if confirmed else None,
    )
    blocked = []
    if current is None:
        blocked.append("研究问题尚未人工确认")
    if not formal:
        blocked.append("没有通过当前动态证据核验的正式科研决策")
    if not run or run["run_id"] not in linked:
        blocked.append("正式决定尚未关联当前修订的已确认实验")
    if not any(
        e["status"] == "complete" and e.get("current_revision")
        for e in snapshot["evaluations"]
    ):
        blocked.append("六题成对评价或独立人工复核尚未完成")
    if snapshot["privacy"]["redacted_fields"]:
        blocked.append("存在敏感字段，须先修正来源记录并重新人工检查")
    result = "实验未完成确认，尚不能报告保留集指标。"
    if run and run.get("test_aggregates"):
        group = run["test_aggregates"][0]
        metrics = group.get("metrics", {})
        f1, recall, fp = (
            metrics.get(key, {})
            for key in ("spam_f1", "spam_recall", "ham_false_positive_rate")
        )
        numbers = [
            f1.get("mean"),
            f1.get("sample_std"),
            recall.get("mean"),
            fp.get("mean"),
        ]
        if all(type(n) in (int, float) and math.isfinite(n) for n in numbers):
            method = {"word_lr": "逻辑回归", "word_nb": "朴素贝叶斯"}.get(
                group.get("candidate"), "所选方法"
            )
            result = f"经哈希核验的实验中，{method}保留集spam F1为{numbers[0]:.4f}±{numbers[1]:.4f}（{f1.get('n', '未记录')}个种子的均值与样本标准差），召回率{numbers[2]:.4f}，正常短信误报率{numbers[3]:.4f}。"
    review = (
        "当前证据及科研决定仍待人工完成核验，不将模型建议当作正式结论。"
        if not formal
        else "正式决定仅引用当前修订下书目、页码、直接引文及人工支持判断均合格的证据。"
    )
    paired = (
        "成对评价尚待完成，不声称效率提升或正确率优势。"
        if not any(
            e["status"] == "complete" and e.get("current_revision")
            for e in snapshot["evaluations"]
        )
        else "成对评价已记录独立人工判断与耗时，仅报告本案例观察，不外推普遍优势。"
    )
    paragraphs = [
        "问题场景：面对有限训练标注下的英文垃圾短信识别，我们关注TF-IDF搭配朴素贝叶斯或逻辑回归的基线选择。通用AI回答可能混淆文献结论与任务迁移，因此建设本地网页工作台“研证”，把问题、来源、证据、实验和人工决定串成可追溯记录。",
        "解决过程：通过正式来源检索书目，导入合法取得的论文，以PDF物理页码定位引文。AI辅助澄清问题、提取证据及比较预设候选；程序核对原句是否出现，人再判断原文是否支持主张及适用边界。编辑上游记录会使相关复核或决定失效。短信正文只在本地处理，不发送给模型；固定去重、划分和训练预算，先比较验证集，再人工冻结候选后评估保留集。",
        "核验与结果：" + result + review + paired,
        "局限反思：引文命中不等于推论正确，公开可读也不等于可再分发。两篇通用文本分类论文不是本项目SMS管道的原样复现。低召回不能被零误报掩盖。实验仅限历史英文数据、预设方法及当前预算，不能外推到中文反诈或生产环境；不靠模型自评正确率。系统本地运行但账号后端可能联网，模型输出仍需人负责核验。",
    ]
    models = sorted(
        {
            r["actual_model"]
            for r in snapshot["model_runs"]
            if r.get("status") == "success" and r.get("actual_model")
        }
    )
    model_text = "、".join(models) if models else "尚无成功模型调用记录"
    notes = [
        "AI使用声明：" + model_text + "；用于结构化辅助，不替代人审或真实实验。",
        "参考：Wang与Manning（2012），Baselines and Bigrams，ACL P12-2018。",
        "参考：Joulin等（2017），Bag of Tricks，ACL E17-2068。",
    ]
    if (
        character_count("".join(paragraphs)) > 900
        or character_count(TITLE + "".join(paragraphs + notes) + PREVIEW) > 1000
    ):
        raise ExportError(
            "正文或包含标题、声明、参考文献的总字数超限；请缩短实际模型说明"
        )
    return paragraphs, notes, blocked


def _docx(paragraphs: list[str], notes: list[str]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = (
        section.right_margin
    ) = Mm(25)
    # 删除core properties中的默认作者、最后编辑者、模板日期等全部字段。
    for element in list(document.core_properties._element):
        document.core_properties._element.remove(element)
    style = document.styles["Normal"]
    style.font.name, style.font.size = "SimSun", Pt(14)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing = Pt(21)
    style.paragraph_format.first_line_indent = Pt(28)
    style.paragraph_format.space_after = Pt(0)

    def add(text, title=False):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = Pt(21)
        paragraph.paragraph_format.first_line_indent = Pt(0 if title else 28)
        if title:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        run.font.name, run.font.size = (
            ("SimHei", Pt(16)) if title else ("SimSun", Pt(14))
        )
        rfonts = run._element.get_or_add_rPr().rFonts
        rfonts.set(qn("w:eastAsia"), "黑体" if title else "宋体")

    add(TITLE, True)
    add(PREVIEW + "；须在合法安装宋体、黑体的环境导出并复核。")
    for text in paragraphs + notes:
        add(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _case_payload(store: Store):
    snapshot = process_snapshot(store)
    paragraphs, notes, blocked = _case_content(snapshot)
    payload = {
        "title": TITLE,
        "paragraphs": paragraphs,
        "notes": notes,
        "process_sha256": _digest(_bytes(snapshot)),
        "blocked": blocked,
    }
    return payload, snapshot


def record_case_review(
    store: Store,
    output_dir="artifacts",
    *,
    content_reviewed: bool,
    anonymity_confirmed: bool,
) -> dict:
    """仅由作品负责人显式确认调用；绑定全文和最新研究状态，不自动代签。"""
    if content_reviewed is not True or anonymity_confirmed is not True:
        raise ExportError("必须人工确认全文事实、匿名与AI声明")
    root = _directory(output_dir, draft=True)
    payload, _ = _case_payload(store)
    if payload["blocked"]:
        raise ExportError("科研门禁未通过，不能确认正式说明")
    record = {
        "content_sha256": _digest(_bytes(payload)),
        "content_reviewed": True,
        "anonymity_confirmed": True,
        "confirmed_at": now(),
    }
    _write(root / "case-review.json", _bytes(record))
    return record


def import_case_pdf(
    store: Store, pdf_path, output_dir="artifacts", *, fonts_licensed: bool = False
) -> dict:
    """登记在合法Word环境导出的正式PDF；正文、真字体及最新人审均须匹配。"""
    from pypdf import PdfReader

    if fonts_licensed is not True:
        raise ExportError("必须人工确认PDF所用字体的合法使用权")
    root = _directory(output_dir, draft=True)
    payload, _ = _case_payload(store)
    if payload["blocked"]:
        raise ExportError("科研门禁未通过，不能登记正式PDF")
    try:
        review = json.loads(
            _safe_path(root / "case-review.json").read_text(encoding="utf-8")
        )
        if (
            not isinstance(review, dict)
            or review.get("content_sha256") != _digest(_bytes(payload))
            or review.get("content_reviewed") is not True
            or review.get("anonymity_confirmed") is not True
            or not review.get("confirmed_at")
        ):
            raise ExportError("当前说明尚未由作品负责人确认")
        data = _safe_path(pdf_path).read_bytes()
    except (OSError, ValueError):
        raise ExportError("PDF或当前说明人审记录无法安全读取") from None
    checked = _check_pdf_bytes(data)
    if not checked["ok"]:
        raise ExportError("正式PDF文字、字体、匿名元数据或版式技术检查未通过")
    text = "".join(page.extract_text() or "" for page in PdfReader(BytesIO(data)).pages)
    expected = payload["title"] + "".join(payload["paragraphs"] + payload["notes"])
    if re.sub(r"\s", "", text) != re.sub(r"\s", "", expected):
        raise ExportError("PDF全文与当前已人工确认的说明不一致")
    submission = _directory(root / "submission")
    _write(submission / "case.pdf", data, no_replace=True)
    _write(
        submission / "case-provenance.json",
        _bytes(
            {
                "pdf_sha256": _digest(data),
                "content_sha256": _digest(_bytes(payload)),
                "process_sha256": payload["process_sha256"],
                "fonts_licensed": True,
            }
        ),
    )
    return {
        "formal_ready": True,
        "files": ["submission/case.pdf"],
        "human_visual_review_required": True,
    }


def export_case(
    store: Store,
    output_dir="artifacts",
    song_font=None,
    hei_font=None,
    song_index=0,
    hei_index=0,
    *,
    fonts_licensed: bool = False,
) -> dict:
    root = _directory(output_dir, draft=True)
    payload, _ = _case_payload(store)
    paragraphs, notes = payload["paragraphs"], payload["notes"]
    blocked = list(payload["blocked"])
    _write(root / "case.docx", _docx(paragraphs, notes))
    _write(root / "case-content.json", _bytes(payload))
    sections = [("", p) for p in paragraphs + notes]
    files = ["case.docx", "case-content.json"]
    try:
        _write(root / "case-preview.pdf", _pdf(TITLE, sections))
        files.append("case-preview.pdf")
    except ExportError as exc:
        blocked.append(str(exc))
    font_checks = {}
    if fonts_licensed is not True:
        blocked.append(
            "尚未由作品负责人确认所提供字体的合法使用权；fsType许可不能代替授权"
        )
    for label, path, index, expected in (
        ("song", song_font, song_index, "SimSun"),
        ("hei", hei_font, hei_index, "SimHei"),
    ):
        if path is None:
            blocked.append(f"缺少合法指定字体{expected}；预览字体不可替代")
            continue
        try:
            info = font_preflight(path, index, expected=expected)
            font_checks[label] = info
            blocked.extend(info["reasons"])
        except ExportError as exc:
            blocked.append(str(exc))
    review_file = _safe_path(root / "case-review.json")
    review = {}
    if review_file.is_file():
        try:
            review = json.loads(review_file.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            pass
    if (
        not isinstance(review, dict)
        or review.get("content_sha256") != _digest(_bytes(payload))
        or review.get("content_reviewed") is not True
        or review.get("anonymity_confirmed") is not True
        or not review.get("confirmed_at")
    ):
        blocked.append("作品负责人尚未确认当前全文事实与匿名，或确认后记录已变化")
    if not blocked:
        fonts = _fonts(song_font, hei_font, song_index, hei_index)
        data = _pdf(TITLE, sections, preview=False, fonts=fonts)
        check = _check_pdf_bytes(data)
        if check["ok"]:
            submission = _directory(root / "submission")
            _write(submission / "case.pdf", data, no_replace=True)
            _write(
                submission / "case-provenance.json",
                _bytes(
                    {
                        "pdf_sha256": _digest(data),
                        "content_sha256": _digest(_bytes(payload)),
                        "process_sha256": payload["process_sha256"],
                        "fonts_licensed": True,
                    }
                ),
            )
            files.append("submission/case.pdf")
        else:
            blocked.extend(check["reasons"])
    return {
        "status": "blocked" if blocked else "ready_for_visual_review",
        "files": files,
        "paths": {Path(name).name: str(root / name) for name in files},
        "blocked": blocked,
        "blocked_reasons": blocked,
        "font_checks": font_checks,
        "formal_ready": not blocked,
        "title": TITLE,
        "paragraphs": paragraphs,
        "notes": notes,
        "text": "\n\n".join([TITLE, *paragraphs, *notes]),
        "body_characters": character_count("".join(paragraphs)),
        "total_characters": character_count(TITLE + "".join(paragraphs + notes)),
        "human_review_required": True,
        "content_sha256": _digest(_bytes(payload)),
    }


def _check_pdf_bytes(data: bytes) -> dict:
    from pypdf import PdfReader

    reasons = []
    font_records = []
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            raise ValueError
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        count = character_count(text)
        if not text.strip() or count > 1000:
            reasons.append("PDF文字不可提取或总字符数超过1000")
        if TITLE not in re.sub(r"\s", "", text):
            reasons.append("PDF缺少固定案例标题")
        if PREVIEW in text or "非提交版" in text:
            reasons.append("PDF是预览而非正式提交版")
        if SENSITIVE.search(text) or IDENTITY_LABEL.search(text):
            reasons.append("PDF检测到敏感信息或身份标签")
        metadata = reader.metadata or {}
        if any(
            metadata.get(key) not in (None, "")
            for key in ("/Author", "/Creator", "/Subject")
        ):
            reasons.append("PDF身份元数据未清空")
        allowed_metadata = {
            "/Title",
            "/Author",
            "/Creator",
            "/Subject",
            "/Producer",
            "/CreationDate",
            "/ModDate",
            "/Trapped",
        }
        if set(metadata) - allowed_metadata or reader.trailer["/Root"].get("/Metadata"):
            reasons.append("PDF含未审核的自定义或XMP元数据")
        if SENSITIVE.search(str(metadata)) or IDENTITY_LABEL.search(str(metadata)):
            reasons.append("PDF元数据包含敏感信息")
        seen = set()
        for page in reader.pages:
            if (
                abs(float(page.mediabox.width) - 595.276) > 2
                or abs(float(page.mediabox.height) - 841.89) > 2
            ):
                reasons.append("PDF页面不是A4")
            resources = page.get("/Resources", {}).get_object()
            for reference in (
                resources.get("/Font", {}).get_object().values()
                if "/Font" in resources
                else ()
            ):
                font = reference.get_object()
                name = str(font.get("/BaseFont", "")).lstrip("/")
                if name in seen:
                    continue
                seen.add(name)
                descendants = font.get("/DescendantFonts")
                actual = descendants[0].get_object() if descendants else font
                descriptor = actual.get("/FontDescriptor")
                descriptor = descriptor.get_object() if descriptor else {}
                stream = descriptor.get("/FontFile2")
                actual_info = None
                if stream:
                    from fontTools.ttLib import TTFont

                    with TTFont(BytesIO(stream.get_object().get_data())) as embedded:
                        actual_names = {
                            n.toUnicode()
                            for n in embedded["name"].names
                            if n.nameID == 6
                        }
                        fs_type = int(embedded["OS/2"].fsType)
                        actual_info = {
                            "names": sorted(actual_names),
                            "fs_type": fs_type,
                        }
                base = name.split("+")[-1]
                valid = (
                    actual_info is not None
                    and base in {"SimSun", "SimHei"}
                    and any(n.split("+")[-1] == base for n in actual_info["names"])
                    and not actual_info["fs_type"] & (0x2 | 0x100 | 0x200)
                )
                font_records.append(
                    {
                        "name": name,
                        "embedded": bool(stream),
                        "actual": actual_info,
                        "ok": valid,
                    }
                )
                if not valid:
                    reasons.append("PDF存在非指定真实字体、未嵌入字体或禁止嵌入字体")
        if not {"SimSun", "SimHei"}.issubset(
            {f["name"].split("+")[-1] for f in font_records}
        ):
            reasons.append("PDF未同时使用宋体SimSun与黑体SimHei")
        root = reader.trailer["/Root"]
        if any(
            key in root for key in ("/OpenAction", "/AA", "/AcroForm", "/Names")
        ) or any(page.get("/AA") or page.get("/Annots") for page in reader.pages):
            reasons.append("PDF包含未经允许的交互动作、附件或批注")
        return {
            "ok": not reasons,
            "reasons": list(dict.fromkeys(reasons)),
            "characters": count,
            "fonts": font_records,
        }
    except Exception:
        return {
            "ok": False,
            "reasons": ["PDF无法安全解析或字体结构不完整"],
            "characters": None,
            "fonts": font_records,
        }


def _submission_files(path) -> tuple[Path, dict[str, Path]]:
    root = _safe_path(path)
    files = {}
    for name in ("case.pdf", "case.mp4"):
        file = _safe_path(root / name)
        if file.is_file():
            files[name] = file
    return root, files


def _research_gate(store: Store | None, root: Path, hashes: dict) -> list[str]:
    if store is None:
        return ["未提供当前研究存储，无法动态确认最新科研资格，禁止打包"]
    try:
        payload, _ = _case_payload(store)
        if payload["blocked"]:
            return list(payload["blocked"])
        provenance = json.loads(
            _safe_path(root / "case-provenance.json").read_text(encoding="utf-8")
        )
        review = json.loads(
            _safe_path(root.parent / "case-review.json").read_text(encoding="utf-8")
        )
        content_hash = _digest(_bytes(payload))
        if (
            not isinstance(provenance, dict)
            or not isinstance(review, dict)
            or provenance.get("pdf_sha256") != hashes.get("case.pdf")
            or provenance.get("content_sha256") != content_hash
            or provenance.get("process_sha256") != payload["process_sha256"]
            or provenance.get("fonts_licensed") is not True
            or review.get("content_sha256") != content_hash
            or review.get("content_reviewed") is not True
            or review.get("anonymity_confirmed") is not True
            or not review.get("confirmed_at")
        ):
            return ["正式PDF与最新研究状态、合法字体或说明人审绑定不一致"]
    except (ValueError, OSError, KeyError, TypeError):
        return ["正式PDF缺少可核验的当前研究与人工确认记录"]
    return []


def _file_hash(path: Path) -> str:
    with _safe_path(path).open("rb") as handle:
        return hashlib.file_digest(handle, "sha256").hexdigest()


def _copy_verified(source: Path, target: Path, expected: str) -> None:
    """流式复制视频，校验实际复制内容后排他发布，不覆盖已有不同文件。"""
    _safe_path(source)
    _safe_path(target)
    if target.exists():
        if target.is_file() and _file_hash(target) == expected:
            return
        raise ExportError("目标已存在且内容不同，拒绝覆盖")
    fd, temporary = tempfile.mkstemp(prefix=".package-", dir=target.parent)
    try:
        digest = hashlib.sha256()
        with os.fdopen(fd, "wb") as outgoing, source.open("rb") as incoming:
            while block := incoming.read(1024 * 1024):
                digest.update(block)
                outgoing.write(block)
        if digest.hexdigest() != expected:
            raise ExportError("复制期间源文件变化，拒绝打包")
        try:
            os.link(temporary, target)
        except FileExistsError:
            raise ExportError("目标已存在，拒绝覆盖") from None
    finally:
        os.unlink(temporary)


def check_submission(path, *, store: Store | None = None) -> dict:
    from .video import check_video

    root, files = _submission_files(path)
    pdf = (
        _check_pdf_bytes(files["case.pdf"].read_bytes())
        if "case.pdf" in files
        else {"ok": False, "reasons": ["缺少case.pdf"]}
    )
    try:
        video = check_video(root / "case.mp4")
    except ValueError:
        video = {"ok": False, "reasons": ["视频缺失或无法安全检测"]}
    hashes = {name: _file_hash(file) for name, file in files.items()}
    review = {}
    review_file = _safe_path(root / "review.json")
    if review_file.is_file():
        try:
            review = json.loads(review_file.read_text(encoding="utf-8"))
        except (ValueError, OSError):
            pass
    required = (
        "pdf_reviewed",
        "video_reviewed",
        "anonymity_confirmed",
        "claims_confirmed",
    )
    reviewed = (
        isinstance(review, dict)
        and all(review.get(key) is True for key in required)
        and review.get("hashes") == hashes
        and len(hashes) == 2
        and bool(review.get("confirmed_at"))
    )
    technical = pdf["ok"] and video["ok"]
    blocked = list(pdf["reasons"]) + list(video["reasons"])
    research_reasons = _research_gate(store, root, hashes)
    blocked.extend(research_reasons)
    if not reviewed:
        blocked.append(
            "尚未逐页、逐帧完成人工版面、字幕音画、匿名与事实确认，或文件已变化"
        )
    return {
        "ok": technical and reviewed and not research_reasons,
        "technical_ok": technical,
        "research_ok": not research_reasons,
        "formal_ready": technical and reviewed and not research_reasons,
        "pdf": pdf,
        "video": video,
        "human_review_confirmed": reviewed,
        "hashes": hashes,
        "blocked": blocked,
        "blocked_reasons": blocked,
        "notice": "技术检测不替代完整观看、匿名检查或评审系统上传确认",
    }


def record_submission_review(
    path,
    *,
    pdf_reviewed: bool,
    video_reviewed: bool,
    anonymity_confirmed: bool,
    claims_confirmed: bool,
    store: Store | None = None,
) -> dict:
    flags = {
        "pdf_reviewed": pdf_reviewed,
        "video_reviewed": video_reviewed,
        "anonymity_confirmed": anonymity_confirmed,
        "claims_confirmed": claims_confirmed,
    }
    if not all(value is True for value in flags.values()):
        raise ExportError("必须由作品负责人逐项完成人工确认，不可自动代签")
    check = check_submission(path, store=store)
    if not check["technical_ok"] or not check["research_ok"]:
        raise ExportError("技术或当前科研门禁未通过，不能签署提交文件确认")
    record = {**flags, "hashes": check["hashes"], "confirmed_at": now()}
    _write(_safe_path(path) / "review.json", _bytes(record))
    return record


def package_submission(path, school_name: str, *, store: Store | None = None) -> dict:
    """学校名仅参与最终文件名，不写内容、元数据或模型请求。"""
    if not isinstance(school_name, str) or not school_name.strip():
        raise ExportError("请输入学校名，仅用于最终文件名")
    if (
        any(c in school_name for c in "/\\:")
        or ".." in school_name
        or any(unicodedata.category(c).startswith("C") for c in school_name)
    ):
        raise ExportError("学校名不能含路径分隔符、控制字符或双点")
    school_name = school_name.strip()
    stem = school_name + TITLE
    if len((stem + ".pdf").encode("utf-8")) > 255:
        raise ExportError("最终文件名超过文件系统长度限制")
    checked = check_submission(path, store=store)
    if not checked["ok"]:
        raise ExportError("PDF、视频或人工确认尚未通过，禁止打包")
    root, files = _submission_files(path)
    if {name: _file_hash(file) for name, file in files.items()} != checked["hashes"]:
        raise ExportError("检查后文件发生变化，拒绝打包")
    targets = {name: _safe_path(root / (stem + Path(name).suffix)) for name in files}
    for name, target in targets.items():
        if target.exists() and (
            not target.is_file() or _file_hash(target) != checked["hashes"][name]
        ):
            raise ExportError("目标已存在且内容不同，拒绝覆盖")
    for name, target in targets.items():
        _copy_verified(files[name], target, checked["hashes"][name])
    return {
        "ok": True,
        "files": [target.name for target in targets.values()],
        "submitted": False,
        "notice": "已生成本机文件，仍须作品负责人自行上传并确认评审系统接收成功",
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="匿名研究过程与参赛文件导出；不自动完成任何人工确认"
    )
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--check-submission", metavar="DIRECTORY")
    group.add_argument("--process", action="store_true")
    group.add_argument("--case", action="store_true")
    parser.add_argument("--data-dir", default="data")
    parser.add_argument("--output-dir")
    parser.add_argument("--song-font")
    parser.add_argument("--hei-font")
    parser.add_argument("--song-index", type=int, default=0)
    parser.add_argument("--hei-index", type=int, default=0)
    parser.add_argument(
        "--fonts-licensed",
        action="store_true",
        help="本人明确确认所提供字体具有合法使用权",
    )
    args = parser.parse_args()
    try:
        if args.check_submission:
            result = check_submission(args.check_submission, store=Store(args.data_dir))
        elif args.process:
            result = export_process(
                Store(args.data_dir), args.output_dir or "artifacts/process"
            )
        else:
            result = export_case(
                Store(args.data_dir),
                args.output_dir or "artifacts",
                args.song_font,
                args.hei_font,
                args.song_index,
                args.hei_index,
                fonts_licensed=args.fonts_licensed,
            )
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if not result.get("blocked") else 1
    except (ExportError, ValueError):
        print(
            json.dumps(
                {
                    "ok": False,
                    "blocked": [
                        "导出或检查失败，请核对源记录、字体及安全路径；不输出底层敏感信息"
                    ],
                },
                ensure_ascii=False,
            )
        )
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
